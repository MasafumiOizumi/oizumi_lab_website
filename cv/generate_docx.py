"""Render the website-derived CV model using the reference CV's typography."""
import json
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def hyperlink(paragraph, label, url):
    rel = paragraph.part.relate_to(
        url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True
    )
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1155CC')
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.extend([color, underline])
    run.append(props)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def create(data, destination):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    for name, size in [('Normal', 11), ('Title', 26), ('Heading 1', 20), ('Heading 2', 16)]:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.38
        style.paragraph_format.space_after = Pt(0)
    document.core_properties.title = 'Masafumi Oizumi Curriculum Vitae'
    document.core_properties.author = 'Masafumi Oizumi'
    for item in data['paragraphs']:
        kind = item['kind']
        style = {'title': 'Title', 'heading': 'Heading 1', 'subheading': 'Heading 2'}.get(kind, 'Normal')
        paragraph = document.add_paragraph(style=style)
        fmt = paragraph.paragraph_format
        fmt.widow_control = True
        if kind == 'title':
            fmt.space_after = Pt(18)
            fmt.keep_with_next = True
        elif kind in ('heading', 'subheading'):
            fmt.space_before = Pt(20 if kind == 'heading' else 14)
            fmt.space_after = Pt(10)
            fmt.keep_with_next = True
            if item['text'] == 'PUBLICATIONS':
                fmt.page_break_before = True
        elif kind in ('entry_start', 'pub_start', 'pub_title'):
            fmt.keep_with_next = True
        elif kind in ('entry', 'entry_end'):
            fmt.space_after = Pt(10)
        elif kind == 'pub_end':
            fmt.space_after = Pt(12)
        if kind.startswith('pub_'):
            fmt.keep_together = True
            fmt.left_indent = Inches(0.3)
            if kind == 'pub_start':
                fmt.first_line_indent = Inches(-0.3)
        text = item['text']
        cursor = 0
        for link in item['links']:
            paragraph.add_run(text[cursor:link['start']])
            hyperlink(paragraph, link['text'], link['url'])
            cursor = link['start'] + len(link['text'])
        paragraph.add_run(text[cursor:])
    document.save(destination)
    print(f'Created {destination}')


if __name__ == '__main__':
    with open(sys.argv[2], encoding='utf-8') as source:
        create(json.load(source), sys.argv[1])
